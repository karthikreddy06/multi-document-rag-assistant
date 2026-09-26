"""
File Conversion Service Module.
Supports comprehensive bi-directional and cross-format document conversions:
1. CSV <-> XLSX
2. TXT / Markdown -> DOCX / PDF
3. PDF -> DOCX / PPTX
4. DOCX -> PDF / PPTX
5. PPTX -> DOCX / PDF
6. Image (.png, .jpg, .jpeg, .webp) -> PDF
Supports deterministic rules with optional LLM-assisted slide outlining for dense documents.
"""

import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table as ReportLabTable, TableStyle
from reportlab.lib import colors

from app.utils.logger import setup_logger

logger = setup_logger("services.converter")


class ConversionError(Exception):
    """Base error for document conversion failures."""
    pass


class UnsupportedConversionError(ConversionError):
    """Raised when the requested format pair is not supported or missing system tool."""
    pass


# Mapping of supported conversions: source_ext -> set of target_exts
SUPPORTED_CONVERSIONS: Dict[str, Set[str]] = {
    ".csv": {".xlsx"},
    ".xlsx": {".csv"},
    ".txt": {".docx", ".pdf"},
    ".md": {".docx", ".pdf"},
    ".markdown": {".docx", ".pdf"},
    ".pdf": {".docx", ".pptx"},
    ".docx": {".pdf", ".pptx"},
    ".pptx": {".docx", ".pdf"},
    ".png": {".pdf"},
    ".jpg": {".pdf"},
    ".jpeg": {".pdf"},
    ".webp": {".pdf"},
}

MIME_TYPES: Dict[str, str] = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
}


def is_conversion_supported(source_ext: str, target_ext: str) -> bool:
    """Check if conversion from source_ext to target_ext is supported."""
    src = source_ext.lower()
    tgt = target_ext.lower()
    if not src.startswith("."):
        src = f".{src}"
    if not tgt.startswith("."):
        tgt = f".{tgt}"
    return tgt in SUPPORTED_CONVERSIONS.get(src, set())


def get_supported_targets(source_ext: str) -> List[str]:
    """Get list of supported target extensions for a given source extension."""
    src = source_ext.lower()
    if not src.startswith("."):
        src = f".{src}"
    return sorted(list(SUPPORTED_CONVERSIONS.get(src, set())))


def detect_libreoffice() -> Optional[str]:
    """Detect if LibreOffice / soffice is available in the system PATH."""
    for cmd in ("soffice", "libreoffice", "soffice.bin"):
        path = shutil.which(cmd)
        if path:
            return path
    return None


class DocumentConverter:
    """Production service converting documents across business, slide, and image formats."""

    @classmethod
    def convert(
        cls,
        source_bytes: bytes,
        source_filename: str,
        target_format: str,
        llm_generator: Optional[Any] = None,
    ) -> Tuple[bytes, str, str]:
        """
        Convert source_bytes into target_format.

        Args:
            source_bytes: Binary content of the source document.
            source_filename: Original document filename (used to infer source format).
            target_format: Desired target extension (e.g. '.docx', 'pdf', 'pptx').
            llm_generator: Optional LLM generator for intelligent slide summarization.

        Returns:
            Tuple of (converted_bytes, output_filename, mime_type)
        """
        src_path = Path(source_filename)
        src_ext = src_path.suffix.lower()
        tgt_ext = target_format.lower()
        if not tgt_ext.startswith("."):
            tgt_ext = f".{tgt_ext}"

        if not is_conversion_supported(src_ext, tgt_ext):
            supported = get_supported_targets(src_ext)
            raise UnsupportedConversionError(
                f"Conversion from '{src_ext}' to '{tgt_ext}' is not supported. "
                f"Supported targets for '{src_ext}': {supported}"
            )

        output_filename = f"{src_path.stem}{tgt_ext}"
        mime_type = MIME_TYPES.get(tgt_ext, "application/octet-stream")

        try:
            # Excel / CSV
            if src_ext == ".csv" and tgt_ext == ".xlsx":
                out_bytes = cls._csv_to_xlsx(source_bytes)
            elif src_ext == ".xlsx" and tgt_ext == ".csv":
                out_bytes = cls._xlsx_to_csv(source_bytes)

            # TXT & Markdown
            elif src_ext in (".txt",) and tgt_ext == ".docx":
                out_bytes = cls._txt_to_docx(source_bytes)
            elif src_ext in (".md", ".markdown") and tgt_ext == ".docx":
                out_bytes = cls._md_to_docx(source_bytes)
            elif src_ext in (".txt",) and tgt_ext == ".pdf":
                out_bytes = cls._txt_to_pdf(source_bytes)
            elif src_ext in (".md", ".markdown") and tgt_ext == ".pdf":
                out_bytes = cls._md_to_pdf(source_bytes)

            # PDF Conversions
            elif src_ext == ".pdf" and tgt_ext == ".docx":
                out_bytes = cls._pdf_to_docx(source_bytes)
            elif src_ext == ".pdf" and tgt_ext == ".pptx":
                out_bytes = cls._pdf_to_pptx(source_bytes, llm_generator=llm_generator)

            # DOCX Conversions
            elif src_ext == ".docx" and tgt_ext == ".pdf":
                out_bytes = cls._docx_to_pdf(source_bytes)
            elif src_ext == ".docx" and tgt_ext == ".pptx":
                out_bytes = cls._docx_to_pptx(source_bytes, llm_generator=llm_generator)

            # PPTX Conversions
            elif src_ext == ".pptx" and tgt_ext == ".docx":
                out_bytes = cls._pptx_to_docx(source_bytes)
            elif src_ext == ".pptx" and tgt_ext == ".pdf":
                out_bytes = cls._pptx_to_pdf(source_bytes)

            # Image Conversions
            elif src_ext in (".png", ".jpg", ".jpeg", ".webp") and tgt_ext == ".pdf":
                out_bytes = cls._image_to_pdf(source_bytes)

            else:
                raise UnsupportedConversionError(f"Handler not implemented for {src_ext} -> {tgt_ext}")

            logger.info(f"DocumentConverter: Successfully converted {source_filename} to {output_filename} ({len(out_bytes)} bytes)")
            return out_bytes, output_filename, mime_type

        except UnsupportedConversionError:
            raise
        except Exception as e:
            logger.error(f"Conversion error for {source_filename} to {target_format}: {e}")
            raise ConversionError(f"Failed to convert {source_filename} to {target_format}: {e}") from e

    # --------------------------------------------------------------------------
    # CSV / XLSX Handlers
    # --------------------------------------------------------------------------
    @staticmethod
    def _csv_to_xlsx(data: bytes) -> bytes:
        text = data.decode("utf-8", errors="replace")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"

        reader = csv.reader(io.StringIO(text))
        for row in reader:
            ws.append(row)

        out_buf = io.BytesIO()
        wb.save(out_buf)
        return out_buf.getvalue()

    @staticmethod
    def _xlsx_to_csv(data: bytes) -> bytes:
        wb = openpyxl.load_workbook(filename=io.BytesIO(data), data_only=True)
        ws = wb.active
        out_buf = io.StringIO()
        writer = csv.writer(out_buf)

        for row in ws.iter_rows(values_only=True):
            cleaned_row = ["" if cell is None else str(cell) for cell in row]
            if any(cleaned_row):
                writer.writerow(cleaned_row)

        return out_buf.getvalue().encode("utf-8")

    # --------------------------------------------------------------------------
    # TXT / Markdown Handlers
    # --------------------------------------------------------------------------
    @staticmethod
    def _txt_to_docx(data: bytes) -> bytes:
        text = data.decode("utf-8", errors="replace")
        doc = DocxDocument()
        for line in text.splitlines():
            doc.add_paragraph(line)
        out_buf = io.BytesIO()
        doc.save(out_buf)
        return out_buf.getvalue()

    @staticmethod
    def _md_to_docx(data: bytes) -> bytes:
        text = data.decode("utf-8", errors="replace")
        doc = DocxDocument()
        for line in text.splitlines():
            stripped = line.strip()
            if stripped.startswith("### "):
                doc.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                doc.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                doc.add_heading(stripped[2:], level=1)
            elif stripped.startswith(("- ", "* ", "• ")):
                doc.add_paragraph(stripped[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)
        out_buf = io.BytesIO()
        doc.save(out_buf)
        return out_buf.getvalue()

    @staticmethod
    def _txt_to_pdf(data: bytes) -> bytes:
        text = data.decode("utf-8", errors="replace")
        out_buf = io.BytesIO()
        doc = SimpleDocTemplate(out_buf, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
        styles = getSampleStyleSheet()
        normal = styles["Normal"]
        normal.fontSize = 10
        normal.leading = 14

        story = []
        for line in text.splitlines():
            clean = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if not clean.strip():
                story.append(Spacer(1, 10))
            else:
                story.append(Paragraph(clean, normal))

        doc.build(story)
        return out_buf.getvalue()

    @staticmethod
    def _md_to_pdf(data: bytes) -> bytes:
        text = data.decode("utf-8", errors="replace")
        out_buf = io.BytesIO()
        doc = SimpleDocTemplate(out_buf, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
        styles = getSampleStyleSheet()

        h1 = styles["Heading1"]
        h2 = styles["Heading2"]
        h3 = styles["Heading3"]
        body = styles["Normal"]
        bullet = styles["Bullet"]

        story = []
        for line in text.splitlines():
            s = line.strip()
            if not s:
                story.append(Spacer(1, 8))
                continue
            clean = s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            clean = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", clean)

            if clean.startswith("### "):
                story.append(Paragraph(clean[4:], h3))
            elif clean.startswith("## "):
                story.append(Paragraph(clean[3:], h2))
            elif clean.startswith("# "):
                story.append(Paragraph(clean[2:], h1))
            elif clean.startswith(("- ", "* ", "• ")):
                story.append(Paragraph(clean[2:], bullet))
            else:
                story.append(Paragraph(clean, body))

        doc.build(story)
        return out_buf.getvalue()

    # --------------------------------------------------------------------------
    # Image Handlers
    # --------------------------------------------------------------------------
    @staticmethod
    def _image_to_pdf(data: bytes) -> bytes:
        from PIL import Image
        with Image.open(io.BytesIO(data)) as img:
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            out_buf = io.BytesIO()
            img.save(out_buf, format="PDF", resolution=100.0)
            return out_buf.getvalue()

    # --------------------------------------------------------------------------
    # PDF Handlers
    # --------------------------------------------------------------------------
    @staticmethod
    def _pdf_to_docx(data: bytes) -> bytes:
        from pdf2docx import Converter

        with tempfile.TemporaryDirectory() as tmp_dir:
            pdf_path = os.path.join(tmp_dir, "input.pdf")
            docx_path = os.path.join(tmp_dir, "output.docx")
            with open(pdf_path, "wb") as f:
                f.write(data)

            cv = Converter(pdf_path)
            cv.convert(docx_path, start=0, end=None)
            cv.close()

            with open(docx_path, "rb") as f:
                return f.read()

    @classmethod
    def _pdf_to_pptx(cls, data: bytes, llm_generator: Optional[Any] = None) -> bytes:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        page_texts = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                page_texts.append(t.strip())

        full_text = "\n\n".join(page_texts)

        # Use LLM summarization if text is long/dense and LLM is provided
        if llm_generator and len(full_text.split()) > 250:
            llm_slides = cls._summarize_into_slides_with_llm(full_text, llm_generator)
            if llm_slides:
                return cls._build_pptx_from_slide_dicts(llm_slides)

        # Deterministic page-by-page slide extraction
        slides_data = []
        for idx, p_text in enumerate(page_texts, 1):
            lines = [l.strip() for l in p_text.splitlines() if l.strip()]
            if not lines:
                continue
            title = lines[0] if len(lines[0]) < 80 else f"Section {idx}"
            body_lines = lines[1:] if len(lines[0]) < 80 else lines
            bullets = [re.sub(r"^[-*•\d+\.]+\s*", "", l)[:180] for l in body_lines[:6] if l.strip()]
            slides_data.append({"title": title, "bullets": bullets or [lines[0]]})

        if not slides_data:
            slides_data.append({"title": "Presentation", "bullets": ["Document contains no extractable text."]})

        return cls._build_pptx_from_slide_dicts(slides_data)

    # --------------------------------------------------------------------------
    # DOCX Handlers
    # --------------------------------------------------------------------------
    @classmethod
    def _docx_to_pdf(cls, data: bytes) -> bytes:
        soffice = detect_libreoffice()
        if soffice:
            try:
                with tempfile.TemporaryDirectory() as tmp_dir:
                    docx_path = os.path.join(tmp_dir, "input.docx")
                    with open(docx_path, "wb") as f:
                        f.write(data)

                    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, docx_path]
                    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60, check=True)
                    pdf_path = os.path.join(tmp_dir, "input.pdf")
                    if os.path.exists(pdf_path):
                        with open(pdf_path, "rb") as f:
                            return f.read()
            except Exception as e:
                logger.warning(f"LibreOffice DOCX->PDF failed, using pure Python ReportLab fallback: {e}")

        # Deterministic pure Python fallback using python-docx + ReportLab
        return cls._docx_to_pdf_reportlab(data)

    @staticmethod
    def _docx_to_pdf_reportlab(data: bytes) -> bytes:
        doc = DocxDocument(io.BytesIO(data))
        out_buf = io.BytesIO()
        pdf = SimpleDocTemplate(out_buf, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)

        styles = getSampleStyleSheet()
        h1 = styles["Heading1"]
        h2 = styles["Heading2"]
        h3 = styles["Heading3"]
        body = styles["Normal"]
        body.fontSize = 10
        body.leading = 14
        bullet = styles["Bullet"]

        story = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                story.append(Spacer(1, 6))
                continue
            clean = txt.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            s_name = p.style.name.lower() if p.style else ""
            if "heading 1" in s_name or "title" in s_name:
                story.append(Paragraph(clean, h1))
            elif "heading 2" in s_name:
                story.append(Paragraph(clean, h2))
            elif "heading 3" in s_name:
                story.append(Paragraph(clean, h3))
            elif "bullet" in s_name or "list" in s_name:
                story.append(Paragraph(clean, bullet))
            else:
                story.append(Paragraph(clean, body))

        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = [
                    Paragraph(cell.text.strip().replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), body)
                    for cell in row.cells
                ]
                table_data.append(row_data)
            if table_data:
                col_count = len(table_data[0])
                if col_count > 0:
                    avail_width = 504
                    col_width = avail_width / col_count
                    rl_table = ReportLabTable(table_data, colWidths=[col_width] * col_count)
                    rl_table.setStyle(TableStyle([
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]))
                    story.append(Spacer(1, 8))
                    story.append(rl_table)
                    story.append(Spacer(1, 8))

        if not story:
            story.append(Paragraph("Empty document", body))

        pdf.build(story)
        return out_buf.getvalue()

    @classmethod
    def _docx_to_pptx(cls, data: bytes, llm_generator: Optional[Any] = None) -> bytes:
        doc = DocxDocument(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        # Use LLM summarization if dense text
        if llm_generator and len(full_text.split()) > 250:
            llm_slides = cls._summarize_into_slides_with_llm(full_text, llm_generator)
            if llm_slides:
                return cls._build_pptx_from_slide_dicts(llm_slides)

        # Deterministic heading-based slide builder
        slides_data = []
        current_title = "Overview"
        current_bullets = []

        for p in doc.paragraphs:
            txt = p.text.strip()
            if not txt:
                continue
            s_name = p.style.name.lower() if p.style else ""
            is_heading = (
                "heading 1" in s_name or "title" in s_name or
                "heading 2" in s_name or "heading 3" in s_name or
                (txt.startswith("#") and len(txt) < 80)
            )

            if is_heading:
                if current_bullets or current_title != "Overview":
                    slides_data.append({"title": current_title, "bullets": current_bullets})
                current_title = re.sub(r"^#+\s*", "", txt)
                current_bullets = []
            else:
                clean_bullet = re.sub(r"^[-*•\d+\.]+\s*", "", txt)
                if len(clean_bullet) > 200:
                    clean_bullet = clean_bullet[:197] + "..."
                current_bullets.append(clean_bullet)
                if len(current_bullets) >= 5:
                    slides_data.append({"title": current_title, "bullets": current_bullets})
                    current_bullets = []

        if current_bullets or not slides_data:
            slides_data.append({"title": current_title, "bullets": current_bullets or ["Summary points"]})

        return cls._build_pptx_from_slide_dicts(slides_data)

    # --------------------------------------------------------------------------
    # PPTX Handlers
    # --------------------------------------------------------------------------
    @staticmethod
    def _pptx_to_docx(data: bytes) -> bytes:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        doc = DocxDocument()

        for idx, slide in enumerate(prs.slides, 1):
            title_text = f"Slide {idx}"
            if slide.shapes.title and slide.shapes.title.text.strip():
                title_text = f"Slide {idx}: {slide.shapes.title.text.strip()}"
            doc.add_heading(title_text, level=1)

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = p.text.strip()
                        if txt:
                            if p.level > 0:
                                doc.add_paragraph(txt, style="List Bullet")
                            else:
                                doc.add_paragraph(txt)
                elif shape.has_table:
                    table = shape.table
                    rows = len(table.rows)
                    cols = len(table.columns)
                    if rows > 0 and cols > 0:
                        docx_table = doc.add_table(rows=rows, cols=cols)
                        docx_table.style = 'Table Grid'
                        for r_idx, row in enumerate(table.rows):
                            for c_idx, cell in enumerate(row.cells):
                                docx_table.cell(r_idx, c_idx).text = cell.text.strip()
                        doc.add_paragraph("")

            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    p = doc.add_paragraph()
                    run = p.add_run(f"Speaker Notes: {notes}")
                    run.italic = True

        out_buf = io.BytesIO()
        doc.save(out_buf)
        return out_buf.getvalue()

    @classmethod
    def _pptx_to_pdf(cls, data: bytes) -> bytes:
        soffice = detect_libreoffice()
        if soffice:
            try:
                with tempfile.TemporaryDirectory() as tmp_dir:
                    pptx_path = os.path.join(tmp_dir, "input.pptx")
                    with open(pptx_path, "wb") as f:
                        f.write(data)
                    cmd = [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp_dir, pptx_path]
                    subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60, check=True)
                    pdf_path = os.path.join(tmp_dir, "input.pdf")
                    if os.path.exists(pdf_path):
                        with open(pdf_path, "rb") as f:
                            return f.read()
            except Exception as e:
                logger.warning(f"LibreOffice PPTX->PDF failed, falling back to ReportLab: {e}")

        # Pure Python fallback using python-pptx + ReportLab landscape
        return cls._pptx_to_pdf_reportlab(data)

    @staticmethod
    def _pptx_to_pdf_reportlab(data: bytes) -> bytes:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        out_buf = io.BytesIO()
        pdf = SimpleDocTemplate(out_buf, pagesize=landscape(letter), rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "SlideTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=12,
        )
        body_style = ParagraphStyle(
            "SlideBody",
            parent=styles["Normal"],
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
        )
        bullet_style = ParagraphStyle(
            "SlideBullet",
            parent=styles["Bullet"],
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#334155"),
            spaceAfter=4,
        )

        story = []
        for idx, slide in enumerate(prs.slides, 1):
            if idx > 1:
                story.append(PageBreak())

            title_text = f"Slide {idx}"
            if slide.shapes.title and slide.shapes.title.text.strip():
                title_text = slide.shapes.title.text.strip()
            clean_title = title_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(clean_title, title_style))
            story.append(Spacer(1, 10))

            for shape in slide.shapes:
                if shape == slide.shapes.title:
                    continue
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        txt = p.text.strip()
                        if txt:
                            clean = txt.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                            if p.level > 0:
                                story.append(Paragraph(f"• {clean}", bullet_style))
                            else:
                                story.append(Paragraph(clean, body_style))

        if not story:
            story.append(Paragraph("Empty presentation", body_style))

        pdf.build(story)
        return out_buf.getvalue()

    # --------------------------------------------------------------------------
    # Presentation Helpers
    # --------------------------------------------------------------------------
    @staticmethod
    def _build_pptx_from_slide_dicts(slides_data: List[Dict[str, Any]]) -> bytes:
        """Constructs a widescreen 16:9 PowerPoint presentation from slide definitions."""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        for s_info in slides_data:
            title = s_info.get("title", "Slide")
            bullets = s_info.get("bullets", [])

            slide = prs.slides.add_slide(prs.slide_layouts[1])
            if slide.shapes.title:
                slide.shapes.title.text = title
                # Styling
                title_p = slide.shapes.title.text_frame.paragraphs[0]
                title_p.font.size = Pt(32)
                title_p.font.bold = True
                title_p.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

            # Content body
            content_shape = slide.shapes.placeholders[1]
            tf = content_shape.text_frame
            tf.word_wrap = True

            for idx, bullet_text in enumerate(bullets):
                if idx == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = bullet_text
                p.font.size = Pt(18)
                p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                p.space_after = Pt(12)

        out_buf = io.BytesIO()
        prs.save(out_buf)
        return out_buf.getvalue()

    @staticmethod
    def _summarize_into_slides_with_llm(text: str, llm_generator: Any) -> Optional[List[Dict[str, Any]]]:
        """Uses LLM to synthesize a dense document into an executive 3-to-6 slide outline."""
        prompt = (
            "You are an executive presentation designer. Transform the following text into an outline for a "
            "professional 3 to 6 slide PowerPoint presentation.\n\n"
            "Requirements:\n"
            "- Return ONLY a valid JSON array of objects.\n"
            "- Each object must have exactly two keys: 'title' (string, max 8 words) and 'bullets' (list of 3-5 concise bullet strings).\n"
            "- Do not include markdown code block formatting or backticks around the JSON.\n\n"
            f"Document Text:\n{text[:4000]}\n"
        )
        try:
            if hasattr(llm_generator, "_cloud_generate"):
                raw_ans = llm_generator._cloud_generate(prompt, num_predict=600)
            elif hasattr(llm_generator, "generate_answer"):
                raw_ans = llm_generator.generate_answer(prompt, chunks=[])
            else:
                return None

            cleaned = re.sub(r"^```(?:json)?\s*", "", raw_ans.strip(), flags=re.IGNORECASE)
            cleaned = re.sub(r"\s*```$", "", cleaned)
            data = json.loads(cleaned)
            if isinstance(data, list) and len(data) >= 1:
                valid = []
                for item in data:
                    if isinstance(item, dict) and "title" in item and "bullets" in item:
                        valid.append({
                            "title": str(item["title"]),
                            "bullets": [str(b) for b in item["bullets"]]
                        })
                return valid if valid else None
        except Exception as e:
            logger.warning(f"DocumentConverter: LLM slide summarization failed, falling back to deterministic: {e}")
        return None
