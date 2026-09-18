"""
DOCX Parser — extracts text from Word documents using python-docx.
Preserves heading hierarchy, paragraph text, and table cell content.
"""

from pathlib import Path
from typing import List

from app.ingestion.parsers.base import BaseParser
from app.models import Document
from app.utils.logger import setup_logger

logger = setup_logger("ingestion.parsers.docx")


class DOCXParser(BaseParser):
    """Parse .docx files, preserving heading levels, body text, and table data."""

    SUPPORTED_EXTENSIONS: frozenset = frozenset({".docx"})

    def parse(
        self,
        file_path: Path,
        original_filename: str,
        file_hash: str,
    ) -> List[Document]:
        try:
            from docx import Document as DocxDocument  # python-docx
        except ImportError as e:
            raise ImportError(
                "python-docx is required for DOCX parsing. "
                "Install it with: pip install python-docx"
            ) from e

        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        logger.info(f"Parsing DOCX: {original_filename}")

        try:
            doc = DocxDocument(str(file_path))
        except Exception as e:
            raise ValueError(f"Could not open DOCX file '{original_filename}': {e}") from e

        lines: List[str] = []

        # Extract paragraphs with heading markers preserved for section detection
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            # Map Word heading styles to markdown-style headers for the chunker
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    level = 1
                prefix = "#" * min(level, 4)
                lines.append(f"{prefix} {text}")
            else:
                lines.append(text)

        # Extract table content as tab-separated rows
        for table in doc.tables:
            for row in table.rows:
                cell_texts = [cell.text.strip() for cell in row.cells]
                row_text = "\t".join(cell_texts)
                if row_text.strip():
                    lines.append(row_text)

        full_text = "\n".join(lines).strip()
        if not full_text:
            logger.warning(f"No text extracted from DOCX: {original_filename}")
            return []

        meta = self._base_metadata(
            file_path=file_path,
            original_filename=original_filename,
            file_hash=file_hash,
            page_number=1,
            total_pages=1,
            char_count=len(full_text),
            format="docx",
        )

        logger.info(f"DOCXParser: extracted {len(full_text)} chars from '{original_filename}'")
        return [Document(page_content=full_text, metadata=meta)]
