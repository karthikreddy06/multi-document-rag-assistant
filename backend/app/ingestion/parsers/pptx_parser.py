"""
PPTX Parser — extracts text from PowerPoint presentations using python-pptx.
Each slide becomes one Document, preserving title and body text structure.
"""

from pathlib import Path
from typing import List

from app.ingestion.parsers.base import BaseParser
from app.models import Document
from app.utils.logger import setup_logger

logger = setup_logger("ingestion.parsers.pptx")


class PPTXParser(BaseParser):
    """Parse .pptx files slide-by-slide; each slide becomes a Document."""

    SUPPORTED_EXTENSIONS: frozenset = frozenset({".pptx"})

    def parse(
        self,
        file_path: Path,
        original_filename: str,
        file_hash: str,
    ) -> List[Document]:
        try:
            from pptx import Presentation  # python-pptx
        except ImportError as e:
            raise ImportError(
                "python-pptx is required for PPTX parsing. "
                "Install it with: pip install python-pptx"
            ) from e

        if not file_path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        logger.info(f"Parsing PPTX: {original_filename}")

        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            raise ValueError(f"Could not open PPTX file '{original_filename}': {e}") from e

        total_slides = len(prs.slides)
        documents: List[Document] = []

        prs_title = ""
        try:
            if hasattr(prs, "core_properties") and prs.core_properties and prs.core_properties.title:
                prs_title = prs.core_properties.title.strip()
        except Exception:
            prs_title = ""

        for slide_idx, slide in enumerate(prs.slides):
            slide_lines: List[str] = []
            slide_title = ""

            for shape in slide.shapes:
                if shape.has_table:
                    try:
                        table_rows = []
                        for row in shape.table.rows:
                            row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                            if any(row_cells):
                                table_rows.append(" | ".join(row_cells))
                        if table_rows:
                            slide_lines.append("\nTable:\n" + "\n".join(table_rows))
                    except Exception as e:
                        logger.debug(f"Could not extract table from slide {slide_idx + 1}: {e}")

                if not shape.has_text_frame:
                    continue

                # Identify title shapes
                if hasattr(shape, "placeholder_format") and shape.placeholder_format:
                    ph_idx = shape.placeholder_format.idx
                    if ph_idx == 0:  # title placeholder
                        slide_title = shape.text_frame.text.strip()
                        if slide_title:
                            slide_lines.append(f"# {slide_title}")
                        continue

                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text and text != slide_title:
                        slide_lines.append(text)

            # Check slide notes
            try:
                if hasattr(slide, "has_notes_slide") and slide.has_notes_slide and slide.notes_slide:
                    notes_frame = getattr(slide.notes_slide, "notes_text_frame", None)
                    if notes_frame and notes_frame.text:
                        notes_text = notes_frame.text.strip()
                        if notes_text:
                            slide_lines.append(f"Notes: {notes_text}")
            except Exception as e:
                logger.debug(f"Could not extract notes from slide {slide_idx + 1}: {e}")

            slide_text = "\n".join(slide_lines).strip()
            if not slide_text:
                logger.debug(f"Slide {slide_idx + 1} of '{original_filename}' is empty, skipping.")
                continue

            meta = self._base_metadata(
                file_path=file_path,
                original_filename=original_filename,
                file_hash=file_hash,
                page_number=slide_idx + 1,
                total_pages=total_slides,
                char_count=len(slide_text),
                format="pptx",
                slide_title=slide_title,
            )

            documents.append(Document(page_content=slide_text, metadata=meta))

        logger.info(f"PPTXParser: extracted {len(documents)} slides from '{original_filename}'")
        return documents
